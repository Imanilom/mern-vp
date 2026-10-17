import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from polar_baseline_analysis import (
    METRIC_COLUMNS,
    WINDOW_FIELDS,
    build_daily_baselines,
    time_period,
    write_workbook,
)


class PolarBaselineAnalysisTests(unittest.TestCase):
    def make_window(self, index, *, activity="Duduk", period="Pagi"):
        metrics = {metric: float(index + 1) for metric in METRIC_COLUMNS}
        start = datetime(2026, 10, 7, 6, 0) + timedelta(minutes=index * 5)
        return {
            "Participant": "Participant 001",
            "Local Date": "2026-10-07",
            "Time Period": period,
            "Activity": activity,
            "Window Start": start,
            "Window End": start + timedelta(minutes=5),
            "Raw Records": 300,
            "Valid HR Samples": 300,
            "Valid RR Samples": 300,
            "Valid RR Fraction": 1.0,
            "Window Status": "VALID",
            **metrics,
        }

    def test_three_daily_periods(self):
        self.assertEqual(time_period(datetime(2026, 10, 7, 6, 0)), "Pagi")
        self.assertEqual(time_period(datetime(2026, 10, 7, 11, 59)), "Pagi")
        self.assertEqual(time_period(datetime(2026, 10, 7, 12, 0)), "Siang")
        self.assertEqual(time_period(datetime(2026, 10, 7, 17, 59)), "Siang")
        self.assertEqual(time_period(datetime(2026, 10, 7, 18, 0)), "Malam")
        self.assertEqual(time_period(datetime(2026, 10, 7, 5, 59)), "Malam")

    def test_mature_after_one_day_requires_minimum_windows_per_context(self):
        windows = [self.make_window(index) for index in range(20)]
        baselines = build_daily_baselines(windows, min_windows_to_mature=20)
        self.assertEqual(len(baselines), 1)
        self.assertTrue(baselines[0]["One-Day Baseline Mature"])
        self.assertEqual(baselines[0]["Maturity Status"], "MATURE_1_DAY_SIMULATION")
        self.assertEqual(baselines[0]["Valid Windows"], 20)
        self.assertEqual(baselines[0]["Baseline HR Mean (BPM)"], 10.5)

    def test_baseline_does_not_merge_activities_or_time_periods(self):
        windows = [
            self.make_window(index, activity="Duduk")
            for index in range(20)
        ] + [
            self.make_window(index, activity="Berjalan", period="Siang")
            for index in range(20)
        ]
        baselines = build_daily_baselines(windows, min_windows_to_mature=20)
        self.assertEqual(len(baselines), 2)
        self.assertEqual(
            {(row["Activity"], row["Time Period"]) for row in baselines},
            {("Duduk", "Pagi"), ("Berjalan", "Siang")},
        )

    def test_excel_has_readme_windows_and_daily_baseline_sheets(self):
        window = self.make_window(0)
        baseline = build_daily_baselines([window], min_windows_to_mature=1)
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "baseline.xlsx"
            write_workbook(
                output_path,
                [window],
                baseline,
                {"records_received": 300},
                min_windows_to_mature=1,
                timezone_name="Asia/Jakarta",
            )
            workbook = load_workbook(output_path, read_only=True, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Metode", "Window 5 Menit", "Baseline Harian"],
            )
            self.assertEqual(
                list(workbook["Window 5 Menit"].values)[0],
                tuple(WINDOW_FIELDS),
            )
            self.assertEqual(
                list(workbook["Baseline Harian"].values)[1][5],
                "MATURE_1_DAY_SIMULATION",
            )
            workbook.close()


if __name__ == "__main__":
    unittest.main()
