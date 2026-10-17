#!/usr/bin/env python3
"""Analyze raw PolarData into activity/time-specific daily baseline Excel sheets.

This is a read-only simulation utility. It does not modify MongoDB baselines.
Set MONGODB_URI in the process environment; credentials are never printed.
"""

from __future__ import annotations

import argparse
import math
import os
import sys
from collections import defaultdict
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.integrate import trapezoid
from scipy.signal import welch

WINDOW_MINUTES = 5
MIN_VALID_RR_MS = 300.0
MAX_VALID_RR_MS = 2000.0
MIN_RR_SAMPLES_PER_WINDOW = 45
DEFAULT_MIN_WINDOWS_TO_MATURE = 20
MIN_HR_BPM = 30.0
MAX_HR_BPM = 220.0
METRIC_COLUMNS = [
    "HR Mean (BPM)",
    "RMSSD (ms)",
    "SDNN (ms)",
    "DFA α1",
    "LF (ms²)",
    "HF (ms²)",
    "PN50 (%)",
    "DeltaHF (ms²)",
    "HR MIN (BPM)",
    "HR MAX (BPM)",
    "RR MIN (ms)",
    "RR MAX (ms)",
    "RR MEAN (ms)",
]
RAW_PROJECTION = {
    "_id": 1,
    "user_id": 1,
    "timestamp": 1,
    "hr": 1,
    "rr": 1,
    "rrms": 1,
    "activity": 1,
}
WINDOW_FIELDS = [
    "Participant",
    "Local Date",
    "Time Period",
    "Activity",
    "Window Start",
    "Window End",
    "Raw Records",
    "Valid HR Samples",
    "Valid RR Samples",
    "Valid RR Fraction",
    "Window Status",
    *METRIC_COLUMNS,
]
BASELINE_FIELDS = [
    "Participant",
    "Local Date",
    "Time Period",
    "Activity",
    "Valid Windows",
    "Maturity Status",
    "One-Day Baseline Mature",
    "Minimum Windows Required",
]
for _metric in METRIC_COLUMNS:
    BASELINE_FIELDS.extend([f"Baseline {_metric}", f"Baseline SD {_metric}"])


def finite_number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    number = float(value)
    return number if math.isfinite(number) else None


def timestamp_to_datetime(value: Any, tz: ZoneInfo) -> datetime | None:
    if isinstance(value, datetime):
        dt = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return dt.astimezone(tz)
    number = finite_number(value)
    if number is None:
        return None
    # PolarData.timestamp is epoch milliseconds in the API schema; accept seconds
    # as well for old exports.
    epoch_seconds = number / 1000 if abs(number) >= 20_000_000_000 else number
    try:
        return datetime.fromtimestamp(epoch_seconds, tz=tz)
    except (OverflowError, OSError, ValueError):
        return None


def time_period(local_dt: datetime) -> str:
    """Use requested three periods: morning 06-12, afternoon 12-18, night 18-06."""
    hour = local_dt.hour
    if 6 <= hour < 12:
        return "Pagi"
    if 12 <= hour < 18:
        return "Siang"
    return "Malam"


def calculate_dfa_alpha1(rr_ms: list[float]) -> float | None:
    """Port the API's integrated-series, linear-detrended DFA short-range slope."""
    if len(rr_ms) < 10:
        return None
    mean_rr = float(np.mean(rr_ms))
    integrated = np.cumsum(np.asarray(rr_ms, dtype=float) - mean_rr)
    window_sizes = list(range(4, min(16, len(rr_ms) // 2) + 1))
    log_sizes: list[float] = []
    log_fluctuations: list[float] = []

    for size in window_sizes:
        segment_count = len(integrated) // size
        if segment_count == 0:
            continue
        x = np.arange(size, dtype=float)
        x_centered = x - x.mean()
        denominator = float(np.dot(x_centered, x_centered))
        if denominator == 0:
            continue
        squared_residual_sum = 0.0
        for segment_index in range(segment_count):
            segment = integrated[segment_index * size : (segment_index + 1) * size]
            slope = float(np.dot(x_centered, segment - segment.mean()) / denominator)
            intercept = float(segment.mean() - slope * x.mean())
            residual = segment - (slope * x + intercept)
            squared_residual_sum += float(np.mean(residual**2))
        fluctuation = math.sqrt(squared_residual_sum / segment_count)
        if fluctuation > 0 and math.isfinite(fluctuation):
            log_sizes.append(math.log10(size))
            log_fluctuations.append(math.log10(fluctuation))

    if len(log_sizes) < 2:
        return None
    slope = float(np.polyfit(log_sizes, log_fluctuations, 1)[0])
    return slope if math.isfinite(slope) else None


def calculate_frequency_power(rr_ms: list[float]) -> tuple[float | None, float | None]:
    """Estimate LF/HF power from a 4 Hz interpolated NN-interval tachogram."""
    if len(rr_ms) < 45:
        return None, None
    rr = np.asarray(rr_ms, dtype=float)
    beat_times = np.cumsum(rr) / 1000.0
    beat_times -= beat_times[0]
    if beat_times[-1] < 240 or np.any(np.diff(beat_times) <= 0):
        return None, None

    sample_rate = 4.0
    uniform_times = np.arange(0, beat_times[-1], 1.0 / sample_rate)
    if len(uniform_times) < 128:
        return None, None
    tachogram = np.interp(uniform_times, beat_times, rr)
    tachogram = tachogram - float(np.mean(tachogram))
    frequencies, power_density = welch(
        tachogram,
        fs=sample_rate,
        window="hann",
        nperseg=min(256, len(tachogram)),
        detrend="constant",
        scaling="density",
    )

    def band_power(low: float, high: float) -> float | None:
        mask = (frequencies >= low) & (frequencies <= high)
        if np.count_nonzero(mask) < 2:
            return None
        power = float(trapezoid(power_density[mask], frequencies[mask]))
        return power if math.isfinite(power) else None

    return band_power(0.04, 0.15), band_power(0.15, 0.40)


def calculate_window_metrics(
    records: list[dict[str, Any]],
) -> dict[str, Any]:
    heart_rates = [
        hr
        for record in records
        if (hr := finite_number(record.get("hr"))) is not None
        and MIN_HR_BPM <= hr <= MAX_HR_BPM
    ]
    rr_intervals = []
    for record in records:
        rr = finite_number(record.get("rrms"))
        if rr is None:
            rr = finite_number(record.get("rr"))
        if rr is not None and MIN_VALID_RR_MS <= rr <= MAX_VALID_RR_MS:
            rr_intervals.append(rr)

    result: dict[str, Any] = {
        "Raw Records": len(records),
        "Valid HR Samples": len(heart_rates),
        "Valid RR Samples": len(rr_intervals),
        "Valid RR Fraction": len(rr_intervals) / len(records) if records else 0.0,
        "Window Status": (
            "VALID"
            if len(rr_intervals) >= MIN_RR_SAMPLES_PER_WINDOW and heart_rates
            else "INSUFFICIENT_VALID_SAMPLES"
        ),
        "HR Mean (BPM)": float(np.mean(heart_rates)) if heart_rates else None,
        "HR MIN (BPM)": min(heart_rates) if heart_rates else None,
        "HR MAX (BPM)": max(heart_rates) if heart_rates else None,
        "RR MIN (ms)": min(rr_intervals) if rr_intervals else None,
        "RR MAX (ms)": max(rr_intervals) if rr_intervals else None,
        "RR MEAN (ms)": float(np.mean(rr_intervals)) if rr_intervals else None,
        "RMSSD (ms)": None,
        "SDNN (ms)": None,
        "DFA α1": None,
        "LF (ms²)": None,
        "HF (ms²)": None,
        "PN50 (%)": None,
        "DeltaHF (ms²)": None,
    }
    if len(rr_intervals) >= 2:
        rr_array = np.asarray(rr_intervals, dtype=float)
        differences = np.diff(rr_array)
        result["RMSSD (ms)"] = float(np.sqrt(np.mean(differences**2)))
        result["SDNN (ms)"] = float(np.std(rr_array, ddof=1))
        result["PN50 (%)"] = float(np.count_nonzero(np.abs(differences) > 50) / len(differences) * 100)
        result["DFA α1"] = calculate_dfa_alpha1(rr_intervals)
    if len(rr_intervals) >= MIN_RR_SAMPLES_PER_WINDOW:
        lf, hf = calculate_frequency_power(rr_intervals)
        result["LF (ms²)"] = lf
        result["HF (ms²)"] = hf
    return result


def build_windows(
    records: list[dict[str, Any]],
    tz: ZoneInfo,
) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, int]]:
    buckets: dict[tuple[str, str, str, str, int], list[dict[str, Any]]] = defaultdict(list)
    user_ids: set[str] = set()
    counts = {
        "records_received": len(records),
        "records_skipped_missing_user_or_timestamp": 0,
        "records_skipped_invalid_timestamp": 0,
    }

    for record in records:
        user_value = record.get("user_id")
        if user_value is None:
            counts["records_skipped_missing_user_or_timestamp"] += 1
            continue
        local_dt = timestamp_to_datetime(record.get("timestamp"), tz)
        if local_dt is None:
            counts["records_skipped_invalid_timestamp"] += 1
            continue
        user_id = str(user_value)
        user_ids.add(user_id)
        activity = str(record.get("activity") or "Unknown").strip() or "Unknown"
        period = time_period(local_dt)
        epoch_ms = int(local_dt.timestamp() * 1000)
        window_epoch_ms = (epoch_ms // (WINDOW_MINUTES * 60_000)) * (WINDOW_MINUTES * 60_000)
        buckets[(user_id, local_dt.date().isoformat(), period, activity, window_epoch_ms)].append(record)

    participant_ids = {
        user_id: f"Participant {index:03d}"
        for index, user_id in enumerate(sorted(user_ids), start=1)
    }
    windows: list[dict[str, Any]] = []
    for (user_id, local_date, period, activity, window_epoch_ms), window_records in sorted(buckets.items()):
        start = datetime.fromtimestamp(window_epoch_ms / 1000, tz=tz)
        metrics = calculate_window_metrics(window_records)
        windows.append({
            "Participant": participant_ids[user_id],
            "Local Date": local_date,
            "Time Period": period,
            "Activity": activity,
            "Window Start": start.replace(tzinfo=None),
            "Window End": (start + timedelta(minutes=WINDOW_MINUTES)).replace(tzinfo=None),
            **metrics,
        })

    add_delta_hf(windows)
    return windows, participant_ids, counts


def add_delta_hf(windows: list[dict[str, Any]]) -> None:
    """DeltaHF = current HF minus prior valid HF in same daily context."""
    contexts: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for window in windows:
        contexts[(
            window["Participant"],
            window["Local Date"],
            window["Time Period"],
            window["Activity"],
        )].append(window)
    for context_windows in contexts.values():
        context_windows.sort(key=lambda window: window["Window Start"])
        previous: dict[str, Any] | None = None
        for window in context_windows:
            hf = window.get("HF (ms²)")
            if (
                previous is not None
                and hf is not None
                and previous.get("HF (ms²)") is not None
                and window["Window Start"] - previous["Window Start"]
                <= timedelta(minutes=10)
            ):
                window["DeltaHF (ms²)"] = hf - previous["HF (ms²)"]
            if hf is not None:
                previous = window


def build_daily_baselines(
    windows: list[dict[str, Any]],
    min_windows_to_mature: int,
) -> list[dict[str, Any]]:
    groups: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for window in windows:
        groups[(
            window["Participant"],
            window["Local Date"],
            window["Time Period"],
            window["Activity"],
        )].append(window)

    baseline_rows = []
    for (participant, local_date, period, activity), group_windows in sorted(groups.items()):
        valid_windows = [
            window for window in group_windows if window["Window Status"] == "VALID"
        ]
        is_mature = len(valid_windows) >= min_windows_to_mature
        row: dict[str, Any] = {
            "Participant": participant,
            "Local Date": local_date,
            "Time Period": period,
            "Activity": activity,
            "Valid Windows": len(valid_windows),
            "Maturity Status": (
                "MATURE_1_DAY_SIMULATION"
                if is_mature
                else f"LEARNING ({len(valid_windows)}/{min_windows_to_mature} valid windows)"
            ),
            "One-Day Baseline Mature": is_mature,
            "Minimum Windows Required": min_windows_to_mature,
        }
        for metric in METRIC_COLUMNS:
            values = [
                float(value)
                for window in valid_windows
                if (value := finite_number(window.get(metric))) is not None
            ]
            row[f"Baseline {metric}"] = float(np.mean(values)) if values else None
            row[f"Baseline SD {metric}"] = (
                float(np.std(values, ddof=1)) if len(values) >= 2 else None
            )
        baseline_rows.append(row)
    return baseline_rows


def write_sheet(workbook: Workbook, title: str, fields: list[str], rows: list[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append(fields)
    header_fill = PatternFill("solid", fgColor="164E63")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{max(1, len(rows) + 1)}"
    for row in rows:
        worksheet.append([row.get(field) for field in fields])
    for column_index, field in enumerate(fields, start=1):
        width = min(max(len(field) + 2, 12), 28)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
    worksheet.row_dimensions[1].height = 34


def write_workbook(
    output_path: Path,
    windows: list[dict[str, Any]],
    baselines: list[dict[str, Any]],
    source_info: dict[str, Any],
    min_windows_to_mature: int,
    timezone_name: str,
) -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Metode"
    notes = [
        ["Analisis baseline PolarData - simulasi satu hari"],
        ["Tujuan", "Menghitung fitur Polar per 5 menit, aktivitas, tanggal lokal, dan periode Pagi/Siang/Malam."],
        ["Sumber", "MongoDB collection PolarData (read-only); tidak ada perubahan ke collection baseline."],
        ["Zona waktu", timezone_name],
        ["Batas periode", "Pagi 06:00-11:59; Siang 12:00-17:59; Malam 18:00-05:59."],
        ["Jendela analisis", f"{WINDOW_MINUTES} menit; bucket waktu tetap dan tidak digabung lintas aktivitas/periode."],
        ["Baseline harian", "Rerata dan SD sample-window valid per peserta + tanggal + aktivitas + periode."],
        ["Mature simulasi", f"Satu tanggal lokal dan minimal {min_windows_to_mature} jendela 5 menit valid untuk masing-masing aktivitas/periode."],
        ["Valid window", f"Minimal {MIN_RR_SAMPLES_PER_WINDOW} RR valid (rentang {MIN_VALID_RR_MS:.0f}-{MAX_VALID_RR_MS:.0f} ms) dan minimal satu HR valid."],
        ["HRV", "RMSSD, SDNN sample, pNN50 dihitung dari RR berurutan yang valid di setiap jendela."],
        ["DFA α1", "Port pendekatan implementasi API: integrated RR, detrend linear, skala 4-16."],
        ["LF/HF", "Welch PSD dari RR tachogram interpolasi 4 Hz; LF 0.04-0.15 Hz, HF 0.15-0.40 Hz; satuan ms²."],
        ["DeltaHF", "HF jendela ini dikurangi HF valid sebelumnya pada tanggal/aktivitas/periode sama, jika jeda <=10 menit."],
        ["PN50", "Persentase pasangan RR berurutan dengan beda absolut >50 ms; kolom ditulis PN50 agar sesuai permintaan, lazim disebut pNN50."],
        ["HR/RR MIN/MAX/MEAN", "Nilai dalam jendela 5 menit; HR dibatasi 30-220 bpm, RR dibatasi 300-2000 ms."],
        ["Keterbatasan", "Status MATURE_1_DAY_SIMULATION hanya untuk analisis ini, bukan perubahan status baseline production atau validasi klinis."],
        ["Keterbatasan spektral", "Estimasi LF/HF pada jendela singkat sensitif terhadap artefak dan metode; interpretasikan sebagai analisis teknis, bukan diagnosis."],
        ["Raw documents", source_info.get("records_received", 0)],
        ["Window rows", len(windows)],
        ["Daily baseline rows", len(baselines)],
        ["Valid windows", sum(1 for row in windows if row["Window Status"] == "VALID")],
        ["Mongo write operations", 0],
    ]
    for row in notes:
        readme.append(row)
    readme.column_dimensions["A"].width = 30
    readme.column_dimensions["B"].width = 115
    readme.freeze_panes = "A2"
    for cell in readme[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164E63")
    for row in readme.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    write_sheet(workbook, "Window 5 Menit", WINDOW_FIELDS, windows)
    write_sheet(workbook, "Baseline Harian", BASELINE_FIELDS, baselines)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_date(value: str | None, label: str) -> date | None:
    if value is None:
        return None
    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise argparse.ArgumentTypeError(f"{label} harus berformat YYYY-MM-DD") from error


def query_epoch_bounds(
    start_date: date | None,
    end_date: date | None,
    tz: ZoneInfo,
) -> dict[str, Any]:
    query: dict[str, Any] = {}
    bounds: dict[str, int] = {}
    if start_date is not None:
        local_start = datetime.combine(start_date, time.min, tzinfo=tz)
        bounds["$gte"] = int(local_start.timestamp() * 1000)
    if end_date is not None:
        local_end = datetime.combine(end_date + timedelta(days=1), time.min, tzinfo=tz)
        bounds["$lt"] = int(local_end.timestamp() * 1000)
    if bounds:
        query["timestamp"] = bounds
    return query


def connect_and_read(
    mongo_uri: str,
    database_name: str,
    collection_name: str,
    query: dict[str, Any],
) -> list[dict[str, Any]]:
    from pymongo import MongoClient

    client = MongoClient(mongo_uri, serverSelectionTimeoutMS=10_000)
    try:
        client.admin.command("ping")
        collection = client[database_name][collection_name]
        return list(collection.find(query, RAW_PROJECTION).batch_size(2000))
    finally:
        client.close()


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build per-activity/time-period daily Polar baselines and an Excel workbook."
    )
    parser.add_argument(
        "--database",
        default=os.getenv("MONGODB_DATABASE"),
        help="MongoDB database (defaults to MONGODB_DATABASE or the URI path).",
    )
    parser.add_argument(
        "--collection",
        default="polardatas",
        help="MongoDB collection; Mongoose model PolarData defaults to 'polardatas'.",
    )
    parser.add_argument(
        "--timezone",
        default="Asia/Jakarta",
        help="IANA timezone used for local date and time periods (default: Asia/Jakarta).",
    )
    parser.add_argument("--start-date", help="Inclusive local date YYYY-MM-DD.")
    parser.add_argument("--end-date", help="Inclusive local date YYYY-MM-DD.")
    parser.add_argument(
        "--min-windows-to-mature",
        type=int,
        default=DEFAULT_MIN_WINDOWS_TO_MATURE,
        help=f"Valid 5-minute windows required per day/activity/period (default: {DEFAULT_MIN_WINDOWS_TO_MATURE}).",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(__file__).resolve().parent / "outputs" / "polar_baseline_analysis.xlsx",
        help="Output workbook path.",
    )
    parser.add_argument(
        "--self-test",
        action="store_true",
        help="Run local calculation checks without connecting to MongoDB.",
    )
    return parser.parse_args(argv)


def run_self_test() -> None:
    tz = ZoneInfo("Asia/Jakarta")
    start = datetime(2026, 10, 7, 7, 0, tzinfo=tz)
    records = []
    # A synthetic test fixture only; no readings are exported as real patient data.
    for index in range(300):
        rr = 850 + 25 * math.sin(index / 7) + 8 * math.sin(index / 2)
        records.append({
            "user_id": "synthetic-user",
            "timestamp": int((start + timedelta(seconds=index)).timestamp() * 1000),
            "hr": 60 + 2 * math.sin(index / 30),
            "rr": rr,
            "activity": "Duduk",
        })
    windows, _, _ = build_windows(records, tz)
    assert len(windows) == 1
    assert windows[0]["Window Status"] == "VALID"
    metrics = calculate_window_metrics(records)
    assert metrics["HR Mean (BPM)"] is not None
    assert metrics["RMSSD (ms)"] is not None
    assert metrics["SDNN (ms)"] is not None
    assert metrics["PN50 (%)"] is not None
    assert metrics["DFA α1"] is not None
    # Less than the configured window count must never be marked mature.
    assert build_daily_baselines(windows, 20)[0]["One-Day Baseline Mature"] is False
    assert build_daily_baselines(windows, 1)[0]["One-Day Baseline Mature"] is True
    assert time_period(start) == "Pagi"
    assert time_period(start.replace(hour=14)) == "Siang"
    assert time_period(start.replace(hour=22)) == "Malam"
    print("Self-test passed (synthetic fixture only; MongoDB was not accessed).")


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.self_test:
        run_self_test()
        return 0

    try:
        tz = ZoneInfo(args.timezone)
    except ZoneInfoNotFoundError:
        print(f"Timezone IANA tidak dikenal: {args.timezone}", file=sys.stderr)
        return 2
    if args.min_windows_to_mature < 1:
        print("--min-windows-to-mature harus minimal 1.", file=sys.stderr)
        return 2

    start_date = parse_date(args.start_date, "--start-date")
    end_date = parse_date(args.end_date, "--end-date")
    if start_date and end_date and end_date < start_date:
        print("--end-date tidak boleh sebelum --start-date.", file=sys.stderr)
        return 2

    mongo_uri = os.getenv("MONGODB_URI")
    if not mongo_uri:
        print(
            "MONGODB_URI belum diset. Atur URI MongoDB terautentikasi di environment "
            "terminal lokal; jangan masukkan credential ke source code atau kirim di chat.",
            file=sys.stderr,
        )
        return 2
    database_name = args.database
    if not database_name:
        # Resolve DB from URI path without printing the URI (which may contain secrets).
        from urllib.parse import unquote, urlsplit

        database_name = unquote(urlsplit(mongo_uri).path.lstrip("/").split("/")[0])
    if not database_name:
        print(
            "Database belum ditentukan. Gunakan --database atau MONGODB_DATABASE.",
            file=sys.stderr,
        )
        return 2

    query = query_epoch_bounds(start_date, end_date, tz)
    try:
        records = connect_and_read(mongo_uri, database_name, args.collection, query)
    except Exception as error:
        print(
            f"Gagal membaca MongoDB database/collection yang diminta: {type(error).__name__}: {error}",
            file=sys.stderr,
        )
        return 1

    windows, _, counts = build_windows(records, tz)
    baselines = build_daily_baselines(windows, args.min_windows_to_mature)
    write_workbook(
        args.output,
        windows,
        baselines,
        counts,
        args.min_windows_to_mature,
        args.timezone,
    )
    mature_count = sum(row["One-Day Baseline Mature"] for row in baselines)
    print(f"MongoDB records read: {len(records)} (read-only)")
    print(f"5-minute windows: {len(windows)}; valid: {sum(w['Window Status'] == 'VALID' for w in windows)}")
    print(f"Daily baselines: {len(baselines)}; mature in one-day simulation: {mature_count}")
    print(f"Excel generated: {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
